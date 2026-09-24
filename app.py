import streamlit as st
import requests
import pandas as pd
import plotly.express as px
from io import BytesIO
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------
# 1. Ρυθμίσεις Σελίδας & Design
# ---------------------------------------------------------
st.set_page_config(
    page_title="Metalstorm Stats Scout",
    page_icon="✈️",
    layout="wide"
)

# Custom CSS
st.markdown("""
    <style>
    .stApp {
        background-color: #0f172a;
        color: #f8fafc;
    }
    div[data-testid="stMetric"] {
        background-color: #1e293b;
        padding: 15px;
        border-radius: 10px;
        border: 1px solid #334155;
    }
    div[data-testid="stMetricValue"] {
        font-size: 1.7rem !important;
        color: #38bdf8 !important;
    }
    .main-header {
        font-size: 2.2rem;
        font-weight: 700;
        color: #38bdf8;
        margin-bottom: 0.2rem;
    }
    </style>
""", unsafe_allow_html=True)

st.markdown('<div class="main-header">✈️ Metalstorm Scout — Lifetime Stats</div>', unsafe_allow_html=True)
st.caption("Πλήρης ανάλυση προφίλ, αεροσκαφών, Game Modes & μορφοποιημένη εξαγωγή Excel")

# ---------------------------------------------------------
# 2. Sidebar - Εισαγωγή Startoken & Οδηγίες
# ---------------------------------------------------------
with st.sidebar:
    st.header("🔑 Σύνδεση")
    token_input = st.text_input("Εισάγετε το Startoken σας:", type="password", help="Επικολλήστε το token")
    fetch_btn = st.button("🔄 Φόρτωση Στατιστικών", type="primary", use_container_width=True)
    
    st.divider()
    st.markdown("### ❓ Πώς να βρείτε το Token σας")
    
    tab_pc, tab_mobile = st.tabs(["💻 Υπολογιστής", "📱 Κινητό"])
    
    with tab_pc:
        st.markdown("""
        1. Ανοίξτε το **playmetalstorm.com/stats** στον Chrome/Edge.
        2. Πατήστε **F12** ➔ καρτέλα **Network**.
        3. Επιλέξτε το φίλτρο **Fetch/XHR**.
        4. Βρείτε το αίτημα **`my`**.
        5. Στο **Headers ➔ Authorization**, αντιγράψτε το κείμενο μετά το `startoken `.
        """)
        
    with tab_mobile:
        st.markdown("""
        **Android (Kiwi Browser):**
        1. Ανοίξτε το **playmetalstorm.com/stats** στο Kiwi Browser.
        2. Πατήστε 3 τελείες ➔ **Developer Tools ➔ Network**.
        3. Βρείτε το `my` και αντιγράψτε το token.

        **iPhone (iOS):**
        * Χρησιμοποιήστε εφαρμογή όπως το **Web Inspector** για να δείτε τα Network Headers.
        """)

# ---------------------------------------------------------
# 3. Συνάρτηση Ανάκτησης Δεδομένων
# ---------------------------------------------------------
def fetch_user_stats(token):
    url = "https://playmetalstorm.com/stats/users/my"
    clean_token = token.replace("startoken ", "").strip()
    
    headers = {
        "Authorization": f"startoken {clean_token}",
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
        "Accept": "application/json"
    }
    
    try:
        response = requests.get(url, headers=headers, timeout=10)
        if response.status_code == 200:
            return response.json(), None
        elif response.status_code == 401:
            return None, "Μη έγκυρο ή ληγμένο Token."
        else:
            return None, f"Σφάλμα διακομιστή ({response.status_code})."
    except Exception as e:
        return None, f"Σφάλμα σύνδεσης: {str(e)}"

# ---------------------------------------------------------
# 4. Συνάρτηση Δημιουργίας Styled Excel
# ---------------------------------------------------------
def generate_styled_excel(df_profile, df_planes, df_modes):
    output = BytesIO()
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Αφαίρεση default sheet
    
    # Στυλ
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Σκούρο Μπλε
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    sheets_data = [
        ('Προφίλ & Σύνολα', df_profile),
        ('Στατιστικά Αεροσκαφών', df_planes),
        ('Game Modes', df_modes)
    ]
    
    for sheet_name, df in sheets_data:
        ws = wb.create_sheet(title=sheet_name)
        
        # Headers
        headers = list(df.columns)
        ws.append(headers)
        
        # Format Headers
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        # Data
        for row in df.itertuples(index=False):
            ws.append(list(row))
            
        # Format Data Rows
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.font = data_font
                cell.border = thin_border
                if isinstance(cell.value, float):
                    cell.number_format = '0.0'
                elif isinstance(cell.value, int):
                    cell.number_format = '#,##0'

        # Auto adjust column width
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    wb.save(output)
    return output.getvalue()

# Execution Event
if fetch_btn and token_input:
    with st.spinner("Ανάκτηση δεδομένων..."):
        data, error = fetch_user_stats(token_input)
        if error:
            st.error(error)
        else:
            st.session_state['raw_data'] = data

# ---------------------------------------------------------
# 5. Επεξεργασία & Προβολή Στατιστικών
# ---------------------------------------------------------
if 'raw_data' in st.session_state:
    raw = st.session_state['raw_data']
    
    try:
        accounts = raw.get("accounts", {})
        first_account_id = list(accounts.keys())[0]
        account_obj = accounts[first_account_id]
        player_data = account_obj.get("playerData", {})
    except Exception:
        player_data = raw

    # Profile Attributes
    player_name = player_data.get("playerName", "Unknown")
    friend_code = player_data.get("friendCode", "-")
    trophies = player_data.get("cachedTotalVehicleRankingPoints", 0)
    peak_trophies = player_data.get("highestTotalVehicleRankingPoints", 0)
    player_xp = player_data.get("playerXp", 0)
    hangar_points = player_data.get("cachedHangerPointsSum", 0)
    equipped_plane_raw = player_data.get("equippedVehicleConfig", "-")
    equipped_plane = equipped_plane_raw.replace("aircraft:", "").replace("-v2", "").replace("-v3", "").replace("-v4", "").replace("-", " ").title()
    pilot_banner = player_data.get("pilotBanner", "-").replace("pilot-banner:", "").replace("-", " ").title()
    
    # Career Counters
    career_counters = player_data.get("careerCounters", {})
    counters_by_plane = career_counters.get("countersByVehicleConfig", {})
    overall_stats = career_counters.get("overall", {})
    
    total_pvp_matches = overall_stats.get("pvpMatchesPlayed", 0)
    total_pvp_wins = overall_stats.get("pvpMatchesWon", 0)
    total_mvps = overall_stats.get("pvpMvpCount", 0)
    total_deaths = overall_stats.get("pvpDeaths", 0)
    total_score = overall_stats.get("pvpScore", 0)
    
    # ---------------------------------------------------------
    # A. Επεξεργασία Στατιστικών Αεροσκαφών (Lifetime Stats)
    # ---------------------------------------------------------
    planes_list = []

    for plane_id, stats in counters_by_plane.items():
        clean_name = plane_id.replace("aircraft:", "").replace("-v2", "").replace("-v3", "").replace("-v4", "").replace("-", " ").title()
        
        matches = stats.get("pvpMatchesPlayed", 0)
        wins = stats.get("pvpMatchesWon", 0)
        deaths = stats.get("pvpDeaths", 0)
        score = stats.get("pvpScore", 0)
        mvps = stats.get("pvpMvpCount", 0)
        
        win_rate = round((wins / matches) * 100, 1) if matches > 0 else 0.0
        avg_score = round(score / matches, 1) if matches > 0 else 0.0
        deaths_per_match = round(deaths / matches, 2) if matches > 0 else 0.0

        planes_list.append({
            "Αεροσκάφος": clean_name,
            "Αγώνες": matches,
            "Νίκες": wins,
            "Win Rate (%)": win_rate,
            "Θάνατοι": deaths,
            "Deaths/Match": deaths_per_match,
            "MVPs": mvps,
            "Μέσο Σκορ": avg_score
        })

    df_planes = pd.DataFrame(planes_list)
    df_planes_sorted = df_planes.sort_values(by="Αγώνες", ascending=False) if not df_planes.empty else df_planes

    # Overall Ratios
    overall_win_rate = round((total_pvp_wins / total_pvp_matches) * 100, 1) if total_pvp_matches > 0 else 0.0
    overall_d_m = round(total_deaths / total_pvp_matches, 2) if total_pvp_matches > 0 else 0.0

    # ---------------------------------------------------------
    # B. Επεξεργασία Game Modes
    # ---------------------------------------------------------
    game_modes_data = career_counters.get("countersByGameModeCategory", {})
    modes_list = []
    
    for mode_key, mode_stats in game_modes_data.items():
        clean_mode = mode_key.replace("GameModeCategory", "").replace("GameMode", "")
        mode_names_map = {
            "ControlPoints": "Control Points 🚩",
            "Deathmatch": "Deathmatch ⚔️",
            "PriorityTarget": "Priority Target 🎯",
            "CaptureTheFlag": "Capture The Flag 🚩",
            "Elimination": "Elimination 💀"
        }
        display_mode = mode_names_map.get(clean_mode, clean_mode)
        
        m_matches = mode_stats.get("pvpMatchesPlayed", 0)
        m_wins = mode_stats.get("pvpMatchesWon", 0)
        m_mvps = mode_stats.get("pvpMvpCount", 0)
        m_deaths = mode_stats.get("pvpDeaths", 0)
        m_score = mode_stats.get("pvpScore", 0)
        
        m_wr = round((m_wins / m_matches) * 100, 1) if m_matches > 0 else 0.0
        m_avg_score = round(m_score / m_matches, 1) if m_matches > 0 else 0.0
        
        modes_list.append({
            "Game Mode": display_mode,
            "Αγώνες": m_matches,
            "Νίκες": m_wins,
            "Win Rate (%)": m_wr,
            "MVPs": m_mvps,
            "Θάνατοι": m_deaths,
            "Μέσο Σκορ": m_avg_score
        })
        
    df_modes = pd.DataFrame(modes_list).sort_values(by="Αγώνες", ascending=False) if modes_list else pd.DataFrame()

    # ---------------------------------------------------------
    # C. Δημιουργία Dataframe για Excel Export
    # ---------------------------------------------------------
    profile_export = {
        "Όνομα Παίκτη": player_name,
        "Friend Code": friend_code,
        "Τρόπαια": trophies,
        "Peak Τρόπαια": peak_trophies,
        "Player XP": player_xp,
        "Hangar Points": hangar_points,
        "Συνολικό Win Rate (%)": overall_win_rate,
        "Σύνολο Αγώνων": total_pvp_matches,
        "Σύνολο Νικών": total_pvp_wins,
        "Σύνολο Θανάτων": total_deaths,
        "Overall Deaths/Match": overall_d_m,
        "Σύνολο MVPs": total_mvps
    }
    df_profile_export = pd.DataFrame([profile_export])

    # --- HEADER & EXPORT BUTTON ---
    col_title, col_export = st.columns([3, 1])
    with col_title:
        st.subheader(f"👤 Παίκτης: **{player_name}** | Tag: `{friend_code}`")
        st.caption(f"Εξοπλισμένο Αεροσκάφος: **{equipped_plane}** | Banner: `{pilot_banner}`")
    
    with col_export:
        excel_bytes = generate_styled_excel(df_profile_export, df_planes_sorted, df_modes)
        st.download_button(
            label="📥 Εξαγωγή σε Excel (.xlsx)",
            data=excel_bytes,
            file_name=f"Metalstorm_Stats_{player_name}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
            type="primary"
        )
    
    st.divider()

    # --- TOP METRIC CARDS ---
    c1, c2, c3, c4, c5, c6 = st.columns(6)
    c1.metric("Τρόπαια 🏆", f"{trophies:,}")
    c2.metric("Peak Τρόπαια 🔝", f"{peak_trophies:,}")
    c3.metric("Win Rate 🎯", f"{overall_win_rate}%")
    c4.metric("Player XP ⭐", f"{player_xp:,}")
    c5.metric("Hangar Points 🏭", f"{hangar_points:,}")
    c6.metric("MVPs 🎖️", f"{total_mvps:,}")
    
    st.divider()

    # --- MAIN TABS ---
    tab_planes, tab_modes, tab_raw = st.tabs(["✈️ Αεροσκάφη", "🎮 Ανάλυση Game Modes", "📋 Πλήρεις Πίνακες"])

    # ---------------------------------------------------------
    # TAB 1: ΑΕΡΟΣΚΑΦΗ & GRAPHICS
    # ---------------------------------------------------------
    with tab_planes:
        if not df_planes.empty:
            df_filtered = df_planes[df_planes["Αγώνες"] >= 5].sort_values(by="Win Rate (%)", ascending=False)
            
            sub1, sub2 = st.tabs(["📊 Γράφημα Win Rate", "📈 Μέσο Σκορ ανά Αγώνα"])
            
            with sub1:
                st.subheader("Win Rate (%) ανά Αεροσκάφος (min. 5 αγώνες)")
                fig1 = px.bar(
                    df_filtered,
                    x="Αεροσκάφος",
                    y="Win Rate (%)",
                    text="Win Rate (%)",
                    color="Win Rate (%)",
                    color_continuous_scale="Blues",
                    hover_data=["Αγώνες", "Νίκες", "Θάνατοι", "MVPs"]
                )
                fig1.update_traces(texttemplate='%{text}%', textposition='outside')
                fig1.update_layout(template="plotly_dark", xaxis_tickangle=-45, height=450)
                st.plotly_chart(fig1, use_container_width=True)
                
            with sub2:
                st.subheader("Μέσο Σκορ ανά Αγώνα")
                df_score_sorted = df_planes[df_planes["Αγώνες"] >= 5].sort_values(by="Μέσο Σκορ", ascending=False)
                fig2 = px.bar(
                    df_score_sorted,
                    x="Αεροσκάφος",
                    y="Μέσο Σκορ",
                    text="Μέσο Σκορ",
                    color="Μέσο Σκορ",
                    color_continuous_scale="Viridis",
                    hover_data=["Αγώνες", "MVPs"]
                )
                fig2.update_traces(texttemplate='%{text}', textposition='outside')
                fig2.update_layout(template="plotly_dark", xaxis_tickangle=-45, height=450)
                st.plotly_chart(fig2, use_container_width=True)

            st.subheader("📋 Αναλυτικός Πίνακας Στατιστικών Αεροσκαφών")
            st.dataframe(df_planes_sorted, use_container_width=True)

    # ---------------------------------------------------------
    # TAB 2: GAME MODES ANALYSIS
    # ---------------------------------------------------------
    with tab_modes:
        if not df_modes.empty:
            col_gm_chart, col_gm_data = st.columns([3, 2])
            
            with col_gm_chart:
                st.subheader("Αποδοτικότητα (Win Rate %) ανά Game Mode")
                fig_gm = px.bar(
                    df_modes,
                    x="Game Mode",
                    y="Win Rate (%)",
                    text="Win Rate (%)",
                    color="Win Rate (%)",
                    color_continuous_scale="Tealgrn",
                    hover_data=["Αγώνες", "Νίκες", "MVPs"]
                )
                fig_gm.update_traces(texttemplate='%{text}%', textposition='outside')
                fig_gm.update_layout(template="plotly_dark", height=400)
                st.plotly_chart(fig_gm, use_container_width=True)
                
            with col_gm_data:
                st.subheader("Κατανομή Αγώνων")
                fig_pie = px.pie(
                    df_modes,
                    names="Game Mode",
                    values="Αγώνες",
                    hole=0.4,
                    color_discrete_sequence=px.colors.qualitative.Pastel
                )
                fig_pie.update_layout(template="plotly_dark", height=400)
                st.plotly_chart(fig_pie, use_container_width=True)

            st.subheader("📊 Πίνακας Game Modes")
            st.dataframe(df_modes, use_container_width=True)

    # ---------------------------------------------------------
    # TAB 3: RAW TABLES & EXPORT PREVIEW
    # ---------------------------------------------------------
    with tab_raw:
        st.subheader("Προφίλ Παίκτη")
        st.dataframe(df_profile_export, use_container_width=True)
